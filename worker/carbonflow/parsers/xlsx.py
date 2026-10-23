from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import openpyxl

from .. import limits
from ..errors import LimitExceeded, ParseFailed, UnsupportedFile
from ..sniff import inspect_zip
from .base import ParsedUnit, ParseResult, Span

PARSER_NAME = "openpyxl-cells"


def fmt_cell(value: object) -> str:
    """Text for a cell. Numbers keep their shortest exact representation; nothing is rounded."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, dt.datetime):
        # Excel stores dates as midnight datetimes; show a plain date unless a time is really present.
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat()
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value).strip()


def parse_xlsx(
    data: bytes,
    *,
    max_sheets: int = limits.MAX_XLSX_SHEETS,
    max_cells: int = limits.MAX_XLSX_CELLS,
) -> ParseResult:
    zf = inspect_zip(data)
    names = zf.namelist()
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise UnsupportedFile("workbooks with macros are not accepted")
    if any(n.startswith("xl/externalLinks/") for n in names):
        external = True
    else:
        external = False

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as e:
        raise ParseFailed("workbook is corrupted or unreadable") from e

    try:
        if len(wb.worksheets) > max_sheets:
            raise LimitExceeded(f"workbook has {len(wb.worksheets)} sheets; the limit is {max_sheets}")
        formulas = _formula_cells(data)
        units: list[ParsedUnit] = []
        warnings: list[str] = []
        total_cells = 0

        for ws in wb.worksheets:
            parts: list[str] = []
            spans: list[Span] = []
            pos = 0
            uncached: list[str] = []
            for row in ws.iter_rows():
                row_started = False
                for cell in row:
                    if cell.value is None:
                        if cell.coordinate in formulas.get(ws.title, ()):
                            uncached.append(cell.coordinate)
                        continue
                    total_cells += 1
                    if total_cells > max_cells:
                        raise LimitExceeded(f"workbook has more than {max_cells} filled cells")
                    text = fmt_cell(cell.value)
                    if not text:
                        continue
                    if row_started:
                        parts.append("\t")
                        pos += 1
                    elif parts:
                        parts.append("\n")
                        pos += 1
                    row_started = True
                    c0 = pos
                    parts.append(text)
                    pos += len(text)
                    spans.append(
                        Span(text, c0, pos, {"type": "sheet_cell", "sheet": ws.title, "cell": cell.coordinate})
                    )
            meta: dict = {
                "state": ws.sheet_state,
                "merged": [str(r) for r in ws.merged_cells.ranges],
                "dimensions": ws.dimensions,
            }
            if ws.sheet_state != "visible":
                warnings.append(f"sheet '{ws.title}' is {ws.sheet_state}")
            if uncached:
                meta["formulas_without_value"] = uncached[:20]
                warnings.append(
                    f"sheet '{ws.title}': {len(uncached)} formula cell(s) have no saved value "
                    f"(open and save the file in Excel); first: {uncached[:5]}"
                )
            units.append(ParsedUnit("sheet", ws.title, "".join(parts), spans, meta))
        if external:
            warnings.append("workbook links to external files; values may come from outside this file")
    finally:
        wb.close()

    return ParseResult(units, warnings, PARSER_NAME, f"openpyxl-{openpyxl.__version__}")


def _formula_cells(data: bytes) -> dict[str, set[str]]:
    """Coordinates of formula cells per sheet, so formulas that were never calculated can be flagged."""
    out: dict[str, set[str]] = {}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    except Exception:
        return out
    try:
        for ws in wb.worksheets:
            found: set[str] = set()
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and len(found) < 10_000:
                        found.add(cell.coordinate)
            if found:
                out[ws.title] = found
    finally:
        wb.close()
    return out
